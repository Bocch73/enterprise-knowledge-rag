"""Document upload security, immutable versions, workers, and recovery tests."""

import asyncio
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import BinaryIO, cast
from uuid import UUID

import docx
import pytest
from fastapi.testclient import TestClient

from app.cli import process_ingestion_jobs as worker_cli
from app.core.config import Settings
from app.db.session import Database
from app.infrastructure.storage import LocalObjectStorage, ObjectStorage
from app.modules.ingestion.models import IngestionJob
from app.modules.ingestion.parsers import DocumentParsingError, parse_document
from app.modules.ingestion.worker import process_next_job, recover_stale_jobs
from tests.test_authentication import authorization_header, login_user, register_user


def create_knowledge_base(client: TestClient, token: str, name: str = "Policy Base") -> str:
    response = client.post(
        "/api/v1/knowledge-bases",
        headers=authorization_header(token),
        json={"name": name},
    )
    assert response.status_code == 201, response.text
    return str(response.json()["id"])


def upload_markdown(
    client: TestClient,
    token: str,
    knowledge_base_id: str,
    *,
    content: bytes = b"# Leave Policy\n\nEmployees receive annual leave.",
    filename: str = "leave-policy.md",
) -> dict[str, object]:
    response = client.post(
        f"/api/v1/knowledge-bases/{knowledge_base_id}/documents",
        headers=authorization_header(token),
        files={"file": (filename, content, "text/markdown")},
        data={"title": "Leave Policy"},
    )
    assert response.status_code == 201, response.text
    return cast(dict[str, object], response.json())


def run_one_worker(settings: Settings, storage: ObjectStorage | None = None) -> UUID | None:
    async def run() -> UUID | None:
        database = Database(settings)
        try:
            return await process_next_job(
                database, settings, storage=storage, worker_id="test-worker"
            )
        finally:
            await database.dispose()

    return asyncio.run(run())


def test_markdown_upload_worker_status_and_preview(
    client: TestClient,
    settings: Settings,
) -> None:
    register_user(client)
    token = login_user(client)
    knowledge_base_id = create_knowledge_base(client, token)

    uploaded = upload_markdown(client, token, knowledge_base_id)
    version = cast(dict[str, object], uploaded["version"])
    job = cast(dict[str, object], uploaded["ingestion_job"])

    assert job["status"] == "queued"
    assert job["attempt_count"] == 0
    assert len(str(version["sha256"])) == 64
    assert "storage_key" not in version

    not_ready = client.get(
        f"/api/v1/documents/versions/{version['id']}/preview",
        headers=authorization_header(token),
    )
    assert not_ready.status_code == 404

    processed_job_id = run_one_worker(settings)
    assert str(processed_job_id) == job["id"]

    status_response = client.get(
        f"/api/v1/ingestion-jobs/{job['id']}",
        headers=authorization_header(token),
    )
    assert status_response.status_code == 200
    assert status_response.json()["status"] == "completed"
    assert status_response.json()["progress"] == 100

    preview_response = client.get(
        f"/api/v1/documents/versions/{version['id']}/preview?offset=2&limit=12",
        headers=authorization_header(token),
    )
    assert preview_response.status_code == 200
    preview = preview_response.json()
    assert preview["content"] == "Leave Policy"
    assert preview["offset"] == 2
    assert preview["total_characters"] > preview["returned_characters"]
    assert preview["parser_name"] == "utf-8-markdown"


def test_docx_is_validated_and_parsed(client: TestClient, settings: Settings) -> None:
    register_user(client)
    token = login_user(client)
    knowledge_base_id = create_knowledge_base(client, token)
    document = docx.Document()
    document.add_heading("Expense Policy", level=1)
    document.add_paragraph("Receipts are required for reimbursement.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Limit"
    table.cell(0, 1).text = "1000"
    source = BytesIO()
    document.save(source)

    response = client.post(
        f"/api/v1/knowledge-bases/{knowledge_base_id}/documents",
        headers=authorization_header(token),
        files={
            "file": (
                "expense.docx",
                source.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 201, response.text
    version_id = response.json()["version"]["id"]
    run_one_worker(settings)

    preview = client.get(
        f"/api/v1/documents/versions/{version_id}/preview?limit=1000",
        headers=authorization_header(token),
    )
    assert preview.status_code == 200
    assert "Receipts are required" in preview.json()["content"]
    assert "Limit\t1000" in preview.json()["content"]
    assert preview.json()["parser_name"] == "python-docx"


def test_viewer_can_list_but_cannot_upload_or_preview(
    client: TestClient,
    settings: Settings,
) -> None:
    register_user(client, email="owner@example.com")
    owner_token = login_user(client, email="owner@example.com")
    knowledge_base_id = create_knowledge_base(client, owner_token)
    uploaded = upload_markdown(client, owner_token, knowledge_base_id)
    version = cast(dict[str, object], uploaded["version"])
    job = cast(dict[str, object], uploaded["ingestion_job"])
    run_one_worker(settings)

    register_user(client, email="viewer@example.com")
    viewer_token = login_user(client, email="viewer@example.com")
    grant = client.put(
        f"/api/v1/knowledge-bases/{knowledge_base_id}/members",
        headers=authorization_header(owner_token),
        json={"email": "viewer@example.com", "access_level": "viewer"},
    )
    assert grant.status_code == 200

    listed = client.get(
        f"/api/v1/knowledge-bases/{knowledge_base_id}/documents",
        headers=authorization_header(viewer_token),
    )
    assert listed.status_code == 200
    assert len(listed.json()) == 1

    denied_upload = client.post(
        f"/api/v1/knowledge-bases/{knowledge_base_id}/documents",
        headers=authorization_header(viewer_token),
        files={"file": ("denied.md", b"# denied", "text/markdown")},
    )
    assert denied_upload.status_code == 403
    assert denied_upload.json()["error"]["code"] == "permission_denied"

    denied_preview = client.get(
        f"/api/v1/documents/versions/{version['id']}/preview",
        headers=authorization_header(viewer_token),
    )
    assert denied_preview.status_code == 403
    denied_retry = client.post(
        f"/api/v1/ingestion-jobs/{job['id']}/retry",
        headers=authorization_header(viewer_token),
    )
    assert denied_retry.status_code == 403


def test_unauthorized_user_cannot_discover_document_or_job(client: TestClient) -> None:
    register_user(client, email="owner@example.com")
    owner_token = login_user(client, email="owner@example.com")
    knowledge_base_id = create_knowledge_base(client, owner_token)
    uploaded = upload_markdown(client, owner_token, knowledge_base_id)
    document = cast(dict[str, object], uploaded["document"])
    job = cast(dict[str, object], uploaded["ingestion_job"])

    register_user(client, email="outsider@example.com")
    outsider_token = login_user(client, email="outsider@example.com")
    document_response = client.get(
        f"/api/v1/documents/{document['id']}",
        headers=authorization_header(outsider_token),
    )
    job_response = client.get(
        f"/api/v1/ingestion-jobs/{job['id']}",
        headers=authorization_header(outsider_token),
    )

    assert document_response.status_code == 404
    assert job_response.status_code == 404


def test_upload_rejects_extension_media_structure_filename_and_size(
    client: TestClient,
    settings: Settings,
) -> None:
    register_user(client)
    token = login_user(client)
    knowledge_base_id = create_knowledge_base(client, token)
    endpoint = f"/api/v1/knowledge-bases/{knowledge_base_id}/documents"
    headers = authorization_header(token)

    unsupported = client.post(
        endpoint,
        headers=headers,
        files={"file": ("malware.exe", b"MZ", "application/octet-stream")},
    )
    mismatch = client.post(
        endpoint,
        headers=headers,
        files={"file": ("fake.pdf", b"plain text", "text/plain")},
    )
    unsafe_name = client.post(
        endpoint,
        headers=headers,
        files={"file": ("../escape.md", b"# text", "text/markdown")},
    )
    invalid_docx = client.post(
        endpoint,
        headers=headers,
        files={
            "file": (
                "fake.docx",
                b"PK not a real zip",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    settings.document_upload_max_bytes = 1024
    too_large = client.post(
        endpoint,
        headers=headers,
        files={"file": ("large.md", b"a" * 1025, "text/markdown")},
    )

    assert unsupported.status_code == 415
    assert mismatch.json()["error"]["code"] == "content_type_mismatch"
    assert unsafe_name.json()["error"]["code"] == "unsafe_filename"
    assert invalid_docx.json()["error"]["code"] == "invalid_docx_structure"
    assert too_large.status_code == 413


def test_document_versions_are_immutable_and_duplicate_content_is_rejected(
    client: TestClient,
) -> None:
    register_user(client)
    token = login_user(client)
    knowledge_base_id = create_knowledge_base(client, token)
    first_content = b"# Version One\n\nOriginal policy."
    uploaded = upload_markdown(client, token, knowledge_base_id, content=first_content)
    document = cast(dict[str, object], uploaded["document"])

    duplicate = client.post(
        f"/api/v1/documents/{document['id']}/versions",
        headers=authorization_header(token),
        files={"file": ("copy.md", first_content, "text/markdown")},
    )
    second = client.post(
        f"/api/v1/documents/{document['id']}/versions",
        headers=authorization_header(token),
        files={"file": ("updated.md", b"# Version Two\n\nUpdated policy.", "text/markdown")},
    )
    detail = client.get(
        f"/api/v1/documents/{document['id']}",
        headers=authorization_header(token),
    )

    assert duplicate.status_code == 409
    assert duplicate.json()["error"]["code"] == "duplicate_document_version"
    assert second.status_code == 201
    assert second.json()["version"]["version_number"] == 2
    assert detail.status_code == 200
    assert [version["version_number"] for version in detail.json()["versions"]] == [1, 2]


def test_permanent_parse_failure_is_not_retried(
    client: TestClient,
    settings: Settings,
) -> None:
    register_user(client)
    token = login_user(client)
    knowledge_base_id = create_knowledge_base(client, token)
    response = client.post(
        f"/api/v1/knowledge-bases/{knowledge_base_id}/documents",
        headers=authorization_header(token),
        files={"file": ("broken.pdf", b"%PDF-1.4\nbroken\n%%EOF", "application/pdf")},
    )
    assert response.status_code == 201
    job_id = response.json()["ingestion_job"]["id"]
    run_one_worker(settings)

    failed = client.get(f"/api/v1/ingestion-jobs/{job_id}", headers=authorization_header(token))
    retry = client.post(
        f"/api/v1/ingestion-jobs/{job_id}/retry", headers=authorization_header(token)
    )
    assert failed.json()["status"] == "failed"
    assert failed.json()["retryable"] is False
    assert failed.json()["error_code"] == "pdf_parse_failed"
    assert retry.status_code == 409
    assert retry.json()["error"]["code"] == "ingestion_failure_not_retryable"


class UnavailableStorage:
    """A deterministic substitute that simulates a temporary storage outage."""

    def put_file(self, key: str, source: Path) -> None:
        raise AssertionError("put_file is not used by the worker")

    def open_binary(self, key: str) -> BinaryIO:
        raise OSError("simulated outage")

    def delete(self, key: str) -> None:
        raise AssertionError("delete is not used by the worker")


def test_temporary_failure_allows_one_bounded_manual_retry(
    client: TestClient,
    settings: Settings,
) -> None:
    settings.ingestion_max_attempts = 1
    register_user(client)
    token = login_user(client)
    knowledge_base_id = create_knowledge_base(client, token)
    uploaded = upload_markdown(client, token, knowledge_base_id)
    job = cast(dict[str, object], uploaded["ingestion_job"])

    run_one_worker(settings, UnavailableStorage())
    failed = client.get(f"/api/v1/ingestion-jobs/{job['id']}", headers=authorization_header(token))
    assert failed.json()["status"] == "failed"
    assert failed.json()["retryable"] is True

    retried = client.post(
        f"/api/v1/ingestion-jobs/{job['id']}/retry", headers=authorization_header(token)
    )
    duplicate_retry = client.post(
        f"/api/v1/ingestion-jobs/{job['id']}/retry", headers=authorization_header(token)
    )
    assert retried.status_code == 200
    assert retried.json()["status"] == "queued"
    assert retried.json()["max_attempts"] == 2
    assert duplicate_retry.status_code == 409
    assert duplicate_retry.json()["error"]["code"] == "ingestion_job_not_failed"

    run_one_worker(settings)
    completed = client.get(
        f"/api/v1/ingestion-jobs/{job['id']}", headers=authorization_header(token)
    )
    assert completed.json()["status"] == "completed"


def test_stale_running_job_is_recovered(settings: Settings, client: TestClient) -> None:
    register_user(client)
    token = login_user(client)
    knowledge_base_id = create_knowledge_base(client, token)
    uploaded = upload_markdown(client, token, knowledge_base_id)
    job = cast(dict[str, object], uploaded["ingestion_job"])

    async def abandon_and_recover() -> int:
        database = Database(settings)
        try:
            async with database.session_factory() as session:
                stored_job = await session.get(IngestionJob, UUID(str(job["id"])))
                assert stored_job is not None
                stored_job.status = "running"
                stored_job.stage = "parsing"
                stored_job.attempt_count = 1
                stored_job.locked_at = datetime.now(UTC) - timedelta(hours=1)
                stored_job.worker_id = "crashed-worker"
                await session.commit()
            return await recover_stale_jobs(database, settings)
        finally:
            await database.dispose()

    assert asyncio.run(abandon_and_recover()) == 1
    recovered = client.get(
        f"/api/v1/ingestion-jobs/{job['id']}", headers=authorization_header(token)
    )
    assert recovered.json()["status"] == "queued"
    assert recovered.json()["error_code"] == "worker_lost"

    async def exhaust_stale_job() -> int:
        database = Database(settings)
        try:
            async with database.session_factory() as session:
                stored_job = await session.get(IngestionJob, UUID(str(job["id"])))
                assert stored_job is not None
                stored_job.status = "running"
                stored_job.stage = "parsing"
                stored_job.attempt_count = stored_job.max_attempts
                stored_job.locked_at = datetime.now(UTC) - timedelta(hours=1)
                await session.commit()
            return await recover_stale_jobs(database, settings)
        finally:
            await database.dispose()

    assert asyncio.run(exhaust_stale_job()) == 1
    exhausted = client.get(
        f"/api/v1/ingestion-jobs/{job['id']}", headers=authorization_header(token)
    )
    assert exhausted.json()["status"] == "failed"
    assert exhausted.json()["retryable"] is True


def test_temporary_failure_is_automatically_queued_before_limit(
    client: TestClient,
    settings: Settings,
) -> None:
    register_user(client)
    token = login_user(client)
    knowledge_base_id = create_knowledge_base(client, token)
    uploaded = upload_markdown(client, token, knowledge_base_id)
    job = cast(dict[str, object], uploaded["ingestion_job"])

    run_one_worker(settings, UnavailableStorage())
    waiting = client.get(f"/api/v1/ingestion-jobs/{job['id']}", headers=authorization_header(token))
    assert waiting.json()["status"] == "queued"
    assert waiting.json()["attempt_count"] == 1
    assert waiting.json()["error_code"] == "storage_temporarily_unavailable"
    assert run_one_worker(settings) is None


def test_parser_rejects_empty_oversized_invalid_and_unknown_content(settings: Settings) -> None:
    cases = [
        (BytesIO(b"   "), "text/markdown", "no_extractable_text"),
        (BytesIO(b"\xff"), "text/markdown", "invalid_markdown_encoding"),
        (BytesIO(b"data"), "application/unknown", "unsupported_media_type"),
    ]
    for source, media_type, error_code in cases:
        with pytest.raises(DocumentParsingError) as captured:
            parse_document(source, media_type, settings)
        assert captured.value.code == error_code

    settings.document_max_characters = 1000
    with pytest.raises(DocumentParsingError) as captured:
        parse_document(BytesIO(b"a" * 1001), "text/markdown", settings)
    assert captured.value.code == "extracted_text_too_large"


def test_local_worker_command_processes_a_bounded_batch(
    client: TestClient,
    settings: Settings,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    register_user(client)
    token = login_user(client)
    knowledge_base_id = create_knowledge_base(client, token)
    upload_markdown(client, token, knowledge_base_id)
    monkeypatch.setattr(worker_cli, "get_settings", lambda: settings)
    monkeypatch.setattr("sys.argv", ["process_ingestion_jobs", "--max-jobs", "2"])

    worker_cli.main()

    output = capsys.readouterr().out
    assert "Processed ingestion job" in output
    assert "processed 1 job(s)" in output


def test_local_storage_rejects_escape_and_overwrite(tmp_path: Path) -> None:
    storage = LocalObjectStorage(tmp_path / "objects")
    source = tmp_path / "source.txt"
    source.write_text("immutable", encoding="utf-8")

    storage.put_file("safe/source.txt", source)
    with storage.open_binary("safe/source.txt") as stored:
        assert stored.read() == b"immutable"

    try:
        storage.put_file("safe/source.txt", source)
    except FileExistsError:
        pass
    else:
        raise AssertionError("immutable objects must not be overwritten")

    for unsafe_key in ("../escape", "/absolute", "safe/../escape"):
        try:
            storage.open_binary(unsafe_key)
        except ValueError:
            pass
        else:
            raise AssertionError("unsafe storage keys must be rejected")
