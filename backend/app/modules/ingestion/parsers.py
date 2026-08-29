"""Bounded text extraction for the document types accepted by uploads."""

from __future__ import annotations

from dataclasses import dataclass
from typing import BinaryIO

import docx
import pypdf
from docx.document import Document as DocxDocument
from pypdf import PdfReader

from app.core.config import Settings


class DocumentParsingError(Exception):
    """A permanent source-file error that automated retries cannot repair."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class ParsedResult:
    content: str
    page_count: int | None
    parser_name: str
    parser_version: str


def _require_acceptable_text(content: str, settings: Settings) -> str:
    cleaned = content.strip()
    if not cleaned:
        raise DocumentParsingError(
            "no_extractable_text",
            "The document contains no extractable text; image-only files require OCR",
        )
    if len(cleaned) > settings.document_max_characters:
        raise DocumentParsingError(
            "extracted_text_too_large",
            "The extracted text exceeds the configured safety limit",
        )
    return cleaned


def _parse_pdf(source: BinaryIO, settings: Settings) -> ParsedResult:
    try:
        reader = PdfReader(source, strict=False)
        if reader.is_encrypted:
            raise DocumentParsingError(
                "encrypted_pdf_not_supported",
                "Encrypted PDF files are not supported",
            )
        page_count = len(reader.pages)
        if page_count == 0 or page_count > settings.document_pdf_max_pages:
            raise DocumentParsingError(
                "pdf_page_limit_exceeded",
                "The PDF page count is outside the configured limit",
            )
        pages: list[str] = []
        character_count = 0
        for page in reader.pages:
            text = page.extract_text() or ""
            character_count += len(text)
            if character_count > settings.document_max_characters:
                raise DocumentParsingError(
                    "extracted_text_too_large",
                    "The extracted text exceeds the configured safety limit",
                )
            pages.append(text)
    except DocumentParsingError:
        raise
    except Exception as exception:
        raise DocumentParsingError(
            "pdf_parse_failed",
            "The PDF parser could not read this file",
        ) from exception
    return ParsedResult(
        content=_require_acceptable_text("\n\n".join(pages), settings),
        page_count=page_count,
        parser_name="pypdf",
        parser_version=pypdf.__version__,
    )


def _parse_docx(source: BinaryIO, settings: Settings) -> ParsedResult:
    try:
        document: DocxDocument = docx.Document(source)
        blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                line = "\t".join(cell.text.strip() for cell in row.cells)
                if line.strip():
                    blocks.append(line)
    except Exception as exception:
        raise DocumentParsingError(
            "docx_parse_failed",
            "The DOCX parser could not read this file",
        ) from exception
    return ParsedResult(
        content=_require_acceptable_text("\n\n".join(blocks), settings),
        page_count=None,
        parser_name="python-docx",
        parser_version=docx.__version__,
    )


def _parse_markdown(source: BinaryIO, settings: Settings) -> ParsedResult:
    try:
        content = source.read().decode("utf-8")
    except UnicodeDecodeError as exception:
        raise DocumentParsingError(
            "invalid_markdown_encoding",
            "Markdown files must use UTF-8 encoding",
        ) from exception
    return ParsedResult(
        content=_require_acceptable_text(content, settings),
        page_count=None,
        parser_name="utf-8-markdown",
        parser_version="1",
    )


def parse_document(source: BinaryIO, media_type: str, settings: Settings) -> ParsedResult:
    """Select an allowlisted parser using the canonical stored media type."""
    if media_type == "application/pdf":
        return _parse_pdf(source, settings)
    if media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return _parse_docx(source, settings)
    if media_type == "text/markdown":
        return _parse_markdown(source, settings)
    raise DocumentParsingError("unsupported_media_type", "No parser supports this media type")
